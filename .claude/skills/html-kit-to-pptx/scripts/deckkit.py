# -*- coding: utf-8 -*-
"""
deckkit — build a PPTX in the pixel coordinate system of an HTML slide kit.

Why this exists
---------------
A reference HTML slide deck already encodes every decision that matters: canvas
size, margins, type scale, colors, component geometry. The cheapest way to get a
PowerPoint that looks like the kit is to keep working in the kit's pixel grid and
convert to EMU at the last moment. That is all this module does.

Three pieces:
  Theme   — the numbers you read out of the kit (canvas, padding, fonts, colors)
  Slide   — a flat list of primitives placed in kit pixels
  render_pptx / render_preview — two renderers over the same slide model

The second renderer is the important one. PowerPoint is not installed here and
LibreOffice usually is not either, so there is no way to look at the real file.
render_preview() walks the identical primitive list and writes an SVG, which
cairosvg turns into a PNG you can actually inspect. Because both renderers read
the same coordinates and the text is pre-wrapped by the caller, what you see in
the preview is what lands in the .pptx.

Requires: python-pptx, cairosvg.
"""
from __future__ import annotations

import io
import os
import re

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

EMU_PER_PX = 9525          # 96 dpi: 1 px = 1/96 in = 9525 EMU


# --------------------------------------------------------------------------- theme
class Theme(object):
    """Numbers lifted from the reference kit.

    latin_font / ea_font: Arial (or whatever the kit uses) carries no CJK
    glyphs, so PowerPoint silently substitutes a font of its choosing for
    Korean/Japanese/Chinese runs. Naming the companion face explicitly keeps
    that substitution under your control and makes the preview match.
    """

    def __init__(self, width=1280, height=720, pad_x=88, pad_y=64, runner_h=56,
                 latin_font="Arial", ea_font="Malgun Gothic", colors=None):
        self.width = width
        self.height = height
        self.pad_x = pad_x
        self.pad_y = pad_y
        self.runner_h = runner_h
        self.latin_font = latin_font
        self.ea_font = ea_font
        self.colors = dict(colors or {})

    @property
    def content_w(self):
        return self.width - 2 * self.pad_x

    def c(self, key, default="#000000"):
        return self.colors.get(key, default)


DEFAULT_THEME = Theme()


def E(px):
    """kit pixels -> EMU"""
    return Emu(int(round(px * EMU_PER_PX)))


def px_to_pt(px):
    return px * 0.75


def _rgb(hexs):
    return RGBColor.from_string(hexs.lstrip("#").upper())


def has_cjk(s):
    return any(
        "぀" <= ch <= "ヿ" or "㐀" <= ch <= "鿿"
        or "가" <= ch <= "힣" or "㄰" <= ch <= "㆏"
        for ch in s
    )


def _esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# --------------------------------------------------------------------------- model
class Slide(object):
    """One slide as a flat list of primitives in kit pixels.

    Text is passed in as already-broken lines. PowerPoint's line breaking does
    not match a browser's, so letting it wrap is how decks end up with a stray
    word on its own line or text spilling out of a box. Breaking the lines
    yourself costs a few seconds and removes the whole class of bug — and it is
    what lets the preview be trustworthy.
    """

    def __init__(self, theme=DEFAULT_THEME, bg="#ffffff", runner=None, page=None, dark=False):
        self.theme = theme
        self.bg = bg
        self.runner = runner      # left-hand runner text, or None for no footer
        self.page = page          # right-hand page label
        self.dark = dark          # flips runner + default text colors
        self.items = []
        self.notes = None

    # -- primitives ----------------------------------------------------
    def rect(self, x, y, w, h, fill=None, radius=0, line=None, lw=1):
        self.items.append(("rect", x, y, w, h, fill, radius, line, lw))
        return self

    def line(self, x1, y1, x2, y2, color="#e0e0e0", lw=1):
        self.items.append(("line", x1, y1, x2, y2, color, lw))
        return self

    def circle(self, cx, cy, r, fill):
        return self.rect(cx - r, cy - r, 2 * r, 2 * r, fill, radius=r)

    def text(self, x, y, w, lines, size=20, weight=400, color="#1d1d1f",
             lh=None, align="l", spacing=0.0, caps=False):
        """lines: list of str, or list of [(text, color, weight), ...] runs.

        x/y is the top-left of the text block, not a baseline. spacing is
        letter-spacing in kit pixels (negative tightens, as display type in most
        kits does).
        """
        norm = []
        for ln in lines:
            if isinstance(ln, str):
                norm.append([(ln, color, weight)])
            else:
                norm.append([(t, c or color, wt or weight) for (t, c, wt) in ln])
        self.items.append(("text", x, y, w, norm, size, weight, color,
                           lh or size * 1.4, align, spacing, caps))
        return self

    def svg(self, path, x, y, w, h):
        """Vector graphic. Embedded as SVG with a PNG fallback — see _add_svg."""
        self.items.append(("svg", path, x, y, w, h))
        return self

    def image(self, path, x, y, w, h):
        self.items.append(("image", path, x, y, w, h))
        return self

    # -- composed blocks ----------------------------------------------
    def eyebrow(self, s, y=None, color=None, size=14, spacing=1.6):
        t = self.theme
        return self.text(t.pad_x, y if y is not None else t.pad_y, 700, [s.upper()],
                         size, 600,
                         color or (t.c("accent_on_dark", "#2997ff") if self.dark
                                   else t.c("accent", "#0066cc")),
                         lh=size + 2, spacing=spacing)

    def head(self, title_lines, y=94, size=44, lh=52, color=None, rule=True, spacing=-1.0):
        """Eyebrow-title-rule stack. Returns the y where content can start."""
        t = self.theme
        self.text(t.pad_x, y, t.content_w, title_lines, size, 600,
                  color or ("#ffffff" if self.dark else t.c("ink", "#1d1d1f")),
                  lh=lh, spacing=spacing)
        bottom = y + lh * len(title_lines)
        if rule:
            ry = bottom + 22
            self.rect(t.pad_x, ry, 56, 3,
                      t.c("accent_on_dark", "#2997ff") if self.dark else t.c("accent", "#0066cc"))
            return ry + 3
        return bottom

    def bullets(self, x, y, w, items, size=20, lh=30, gap=22, color=None, dot=None, dot_r=3.5):
        """items: list of line-lists. Returns the y below the last bullet."""
        t = self.theme
        color = color or ("#ffffff" if self.dark else t.c("ink", "#1d1d1f"))
        dot = dot or (t.c("accent_on_dark", "#2997ff") if self.dark else t.c("accent", "#0066cc"))
        cy = y
        for lines in items:
            self.circle(x + dot_r, cy + size * 0.55, dot_r, dot)
            self.text(x + 26, cy, w - 26, lines, size, 400, color, lh=lh)
            cy += lh * len(lines) + gap
        return cy


# --------------------------------------------------------------------------- pptx
def _style_run(run, theme, size, weight, color, spacing):
    f = run.font
    f.size = Pt(px_to_pt(size))
    f.bold = weight >= 600
    f.color.rgb = _rgb(color)
    f.name = theme.latin_font
    rPr = run._r.get_or_add_rPr()
    for tag, face in (("a:ea", theme.ea_font), ("a:cs", theme.latin_font)):
        el = rPr.find(qn(tag))
        if el is None:
            rPr.append(parse_xml(
                '<%s xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                'typeface="%s"/>' % (tag, face)))
        else:
            el.set("typeface", face)
    if spacing:
        # spc is in 1/100 pt
        rPr.set("spc", str(int(round(px_to_pt(spacing) * 100))))


def _add_svg(pptx_slide, path, x, y, w, h, counter):
    """Insert an SVG so PowerPoint draws it as vector art.

    OOXML has no plain "svg picture". The shape carries a normal raster blip
    (which old viewers and thumbnailers use) plus an svgBlip extension pointing
    at the real SVG part, which PowerPoint 2016+ prefers. Both parts must be
    related to the slide; the content type for .svg is registered automatically
    because the part declares it.
    """
    import cairosvg
    svg_bytes = open(path, "rb").read()
    png = cairosvg.svg2png(bytestring=svg_bytes,
                           output_width=int(w * 2), output_height=int(h * 2),
                           background_color="#00000000")
    pic = pptx_slide.shapes.add_picture(io.BytesIO(png), E(x), E(y), E(w), E(h))
    partname = PackURI("/ppt/media/graphic%d.svg" % counter)
    part = Part(partname, "image/svg+xml", pptx_slide.part.package, svg_bytes)
    rId = pptx_slide.part.relate_to(part, RT.IMAGE)
    blip = pic._element.blipFill.find(qn("a:blip"))
    blip.append(parse_xml(
        '<a:extLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}">'
        '<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'r:embed="%s"/></a:ext></a:extLst>' % rId))
    return pic


def render_pptx(slides, out_path, asset_dir=".", theme=None):
    theme = theme or (slides[0].theme if slides else DEFAULT_THEME)
    prs = Presentation()
    prs.slide_width = E(theme.width)
    prs.slide_height = E(theme.height)
    blank = prs.slide_layouts[6]
    counter = [0]

    for sl in slides:
        t = sl.theme
        s = prs.slides.add_slide(blank)
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, E(t.width), E(t.height))
        bg.fill.solid()
        bg.fill.fore_color.rgb = _rgb(sl.bg)
        bg.line.fill.background()
        bg.shadow.inherit = False          # PowerPoint's default preset shadow, off

        for it in sl.items:
            kind = it[0]
            if kind == "rect":
                _, x, y, w, h, fill, radius, lineC, lw = it
                shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
                shp = s.shapes.add_shape(shape, E(x), E(y), E(w), E(h))
                if radius:
                    # adjustment is a fraction of the shorter side, not a radius
                    shp.adjustments[0] = min(0.5, float(radius) / max(1.0, min(w, h)))
                if fill:
                    shp.fill.solid()
                    shp.fill.fore_color.rgb = _rgb(fill)
                else:
                    shp.fill.background()
                if lineC:
                    shp.line.color.rgb = _rgb(lineC)
                    shp.line.width = Pt(px_to_pt(lw))
                else:
                    shp.line.fill.background()
                shp.shadow.inherit = False
            elif kind == "line":
                _, x1, y1, x2, y2, c, lw = it
                cn = s.shapes.add_connector(1, E(x1), E(y1), E(x2), E(y2))
                cn.line.color.rgb = _rgb(c)
                cn.line.width = Pt(px_to_pt(lw))
            elif kind == "text":
                _, x, y, w, lines, size, weight, color, lh, align, spacing, caps = it
                box = s.shapes.add_textbox(E(x), E(y - size * 0.22), E(w),
                                           E(lh * len(lines) + size))
                tf = box.text_frame
                tf.word_wrap = False
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                for i, runs in enumerate(lines):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                                   "r": PP_ALIGN.RIGHT}[align]
                    p.line_spacing = Pt(px_to_pt(lh))
                    p.space_before = Pt(0)
                    p.space_after = Pt(0)
                    for txt, c, wt in runs:
                        r = p.add_run()
                        r.text = txt.upper() if caps else txt
                        _style_run(r, t, size, wt, c, spacing)
            elif kind == "svg":
                _, path, x, y, w, h = it
                counter[0] += 1
                _add_svg(s, os.path.join(asset_dir, path), x, y, w, h, counter[0])
            elif kind == "image":
                _, path, x, y, w, h = it
                s.shapes.add_picture(os.path.join(asset_dir, path), E(x), E(y), E(w), E(h))

        if sl.runner:
            y = t.height - t.runner_h
            cn = s.shapes.add_connector(1, 0, E(y), E(t.width), E(y))
            cn.line.color.rgb = _rgb("#3d3d40" if sl.dark else t.c("hairline", "#e0e0e0"))
            cn.line.width = Pt(0.75)
            for text, xpos, width, alignment, col, wt in (
                    (sl.runner, t.pad_x, 700, PP_ALIGN.LEFT,
                     "#cccccc" if sl.dark else "#7a7a7a", 400),
                    (sl.page or "", t.width - t.pad_x - 200, 200, PP_ALIGN.RIGHT,
                     "#ffffff" if sl.dark else "#333333", 600)):
                box = s.shapes.add_textbox(E(xpos), E(y + 18), E(width), E(24))
                tf = box.text_frame
                tf.word_wrap = False
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                p = tf.paragraphs[0]
                p.alignment = alignment
                r = p.add_run()
                r.text = text
                _style_run(r, t, 13, wt, col, 0)

        if sl.notes:
            s.notes_slide.notes_text_frame.text = sl.notes

    prs.save(out_path)
    return out_path


# --------------------------------------------------------------------------- preview
def render_preview(slide, asset_dir="."):
    """Same primitives, drawn as an SVG you can look at. This is the QA loop."""
    t = slide.theme
    b = ['<svg xmlns="http://www.w3.org/2000/svg" width="%d" height="%d" viewBox="0 0 %d %d">'
         % (t.width, t.height, t.width, t.height),
         '<rect width="%d" height="%d" fill="%s"/>' % (t.width, t.height, slide.bg)]
    for it in slide.items:
        kind = it[0]
        if kind == "rect":
            _, x, y, w, h, fill, radius, lineC, lw = it
            st = ' stroke="%s" stroke-width="%s"' % (lineC, lw) if lineC else ""
            b.append('<rect x="%.1f" y="%.1f" width="%.1f" height="%.1f" rx="%.1f" fill="%s"%s/>'
                     % (x, y, w, h, radius, fill or "none", st))
        elif kind == "line":
            _, x1, y1, x2, y2, c, lw = it
            b.append('<line x1="%.1f" y1="%.1f" x2="%.1f" y2="%.1f" stroke="%s" stroke-width="%s"/>'
                     % (x1, y1, x2, y2, c, lw))
        elif kind == "text":
            _, x, y, w, lines, size, weight, color, lh, align, spacing, caps = it
            anchor = {"l": "start", "c": "middle", "r": "end"}[align]
            tx = x if align == "l" else (x + w / 2.0 if align == "c" else x + w)
            for i, runs in enumerate(lines):
                spans = []
                for txt, c, wt in runs:
                    if caps:
                        txt = txt.upper()
                    fam = (t.ea_font + ", " + t.latin_font) if has_cjk(txt) else t.latin_font
                    spans.append('<tspan fill="%s" font-weight="%d" font-family="%s">%s</tspan>'
                                 % (c, wt, fam, _esc(txt)))
                b.append('<text x="%.1f" y="%.1f" font-size="%.1f" text-anchor="%s" '
                         'letter-spacing="%.2f">%s</text>'
                         % (tx, y + lh * i + size * 0.80, size, anchor, spacing, "".join(spans)))
        elif kind in ("svg", "image"):
            _, path, x, y, w, h = it
            full = os.path.join(asset_dir, path)
            if kind == "svg":
                src = open(full, encoding="utf-8").read()
                m = re.search(r'viewBox="0 0 ([\d.]+) ([\d.]+)"', src)
                sw, sh = (float(m.group(1)), float(m.group(2))) if m else (w, h)
                inner = re.sub(r"^<svg[^>]*>", "", src, count=1).replace("</svg>", "")
                b.append('<g transform="translate(%.1f,%.1f) scale(%.4f,%.4f)">%s</g>'
                         % (x, y, w / sw, h / sh, inner))
            else:
                b.append('<image x="%.1f" y="%.1f" width="%.1f" height="%.1f" href="%s"/>'
                         % (x, y, w, h, full))
    if slide.runner:
        ry = t.height - t.runner_h
        b.append('<line x1="0" y1="%d" x2="%d" y2="%d" stroke="%s" stroke-width="1"/>'
                 % (ry, t.width, ry, "#3d3d40" if slide.dark else t.c("hairline", "#e0e0e0")))
        fam = t.ea_font + ", " + t.latin_font
        b.append('<text x="%d" y="%d" font-size="13" font-family="%s" fill="%s">%s</text>'
                 % (t.pad_x, ry + 34, fam, "#cccccc" if slide.dark else "#7a7a7a",
                    _esc(slide.runner)))
        b.append('<text x="%d" y="%d" font-size="13" font-family="%s" font-weight="600" '
                 'fill="%s" text-anchor="end">%s</text>'
                 % (t.width - t.pad_x, ry + 34, t.latin_font,
                    "#ffffff" if slide.dark else "#333333", _esc(slide.page or "")))
    b.append("</svg>")
    return "".join(b)


def write_previews(slides, out_dir, asset_dir=".", scale=1.0):
    """Write slideNN.png (and .svg) for every slide. Look at these before shipping."""
    import cairosvg
    os.makedirs(out_dir, exist_ok=True)
    paths = []
    for i, sl in enumerate(slides, 1):
        svg = render_preview(sl, asset_dir)
        sp = os.path.join(out_dir, "slide%02d.svg" % i)
        pp = os.path.join(out_dir, "slide%02d.png" % i)
        open(sp, "w", encoding="utf-8").write(svg)
        cairosvg.svg2png(bytestring=svg.encode("utf-8"), write_to=pp,
                         output_width=int(sl.theme.width * scale),
                         output_height=int(sl.theme.height * scale),
                         background_color="#ffffff")
        paths.append(pp)
    return paths
