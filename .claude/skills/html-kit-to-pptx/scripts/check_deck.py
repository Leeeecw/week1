# -*- coding: utf-8 -*-
"""
check_deck.py — structural check + text dump for a generated .pptx

Run this before you call a deck done. It catches the failures that are invisible
in a preview image: shapes pushed off the canvas, SVG parts that were written
but never referenced, and text that silently went missing.

    python check_deck.py deck.pptx [--expect-svg 10]

It does not judge layout. For that, look at the preview PNGs from
deckkit.write_previews().
"""
import sys
import zipfile

from pptx import Presentation

EMU_PER_PX = 9525


def main(path, expect_svg=None):
    prs = Presentation(path)
    W, H = prs.slide_width, prs.slide_height
    slides = list(prs.slides)
    print("canvas      : %.3f x %.3f in  (%d x %d px)"
          % (W / 914400.0, H / 914400.0, W / EMU_PER_PX, H / EMU_PER_PX))
    print("slides      : %d" % len(slides))

    oob = []
    empty = []
    for i, s in enumerate(slides, 1):
        has_text = False
        for sh in s.shapes:
            if sh.left is None:
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                has_text = True
            # a 1px slack keeps rounding noise out of the report
            if (sh.left < -EMU_PER_PX or sh.top < -EMU_PER_PX
                    or sh.left + sh.width > W + EMU_PER_PX
                    or sh.top + sh.height > H + EMU_PER_PX):
                oob.append("  slide %02d: shape at (%.0f, %.0f) size %.0fx%.0f px"
                           % (i, sh.left / EMU_PER_PX, sh.top / EMU_PER_PX,
                              sh.width / EMU_PER_PX, sh.height / EMU_PER_PX))
        if not has_text:
            empty.append(i)

    print("off-canvas  : %d" % len(oob))
    for line in oob:
        print(line)
    if empty:
        print("text-free   : slides %s (fine for a full-bleed image slide, "
              "suspicious otherwise)" % empty)

    with zipfile.ZipFile(path) as z:
        media = sorted(n for n in z.namelist() if n.startswith("ppt/media"))
        svgs = [m for m in media if m.endswith(".svg")]
        ct = z.read("[Content_Types].xml").decode()
        refs = sum(z.read(n).decode().count("svgBlip")
                   for n in z.namelist() if n.startswith("ppt/slides/slide"))
        print("media parts : %d (svg %d)" % (len(media), len(svgs)))
        print("svg content-type overrides: %d" % ct.count("image/svg+xml"))
        print("svgBlip references        : %d" % refs)
        if len(svgs) != refs:
            print("  ! every SVG part must be referenced by exactly one svgBlip;"
                  " a mismatch means PowerPoint will show the PNG fallback instead")
        if expect_svg is not None and len(svgs) != expect_svg:
            print("  ! expected %d SVG graphics, found %d" % (expect_svg, len(svgs)))

    print("\n--- text dump ---")
    for i, s in enumerate(slides, 1):
        parts = []
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                parts.append(" / ".join(p.text for p in sh.text_frame.paragraphs
                                        if p.text.strip()))
        print("[%02d] %s" % (i, " | ".join(parts)))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    exp = None
    if "--expect-svg" in sys.argv:
        exp = int(sys.argv[sys.argv.index("--expect-svg") + 1])
    main(sys.argv[1], exp)
