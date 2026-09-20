# -*- coding: utf-8 -*-
"""
Slide model + two renderers.

Every slide is a flat list of primitives placed in the 1280x720 px coordinate
system of slides-reference-apple.html.  render_pptx() writes real PowerPoint
shapes; render_svg() writes a same-geometry preview used for visual QA (there is
no LibreOffice in this environment, so the preview is how the layout gets
checked before shipping).

Text is pre-wrapped by hand into lines and the PowerPoint boxes have word_wrap
off, so both renderers break lines at exactly the same place.
"""
import io
import os
import re

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
import cairosvg

PX = 9525                      # 1 px @96dpi in EMU
SLIDE_W, SLIDE_H = 1280, 720
PAD_X, PAD_Y = 88, 64
RUNNER_H = 56
LATIN = "Arial"
EA = "Malgun Gothic"           # Arial carries no Hangul; this is the CJK companion


def E(px):
    return Emu(int(round(px * PX)))


def _rgb(hexs):
    return RGBColor.from_string(hexs.lstrip("#").upper())


def _has_hangul(s):
    return any("가" <= ch <= "힣" or "㄰" <= ch <= "㆏" for ch in s)


def _esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


class Slide(object):
    def __init__(self, bg="#ffffff", runner=None, page=None, dark=False):
        self.bg = bg
        self.runner = runner
        self.page = page
        self.dark = dark
        self.items = []
        self.notes = None

    # ---- primitives -------------------------------------------------
    def rect(self, x, y, w, h, fill=None, radius=0, line=None, lw=1):
        self.items.append(("rect", x, y, w, h, fill, radius, line, lw))
        return self

    def line(self, x1, y1, x2, y2, color="#e0e0e0", lw=1):
        self.items.append(("line", x1, y1, x2, y2, color, lw))
        return self

    def circle(self, cx, cy, r, fill):
        self.items.append(("rect", cx - r, cy - r, 2 * r, 2 * r, fill, r, None, 1))
        return self

    def text(self, x, y, w, lines, size=20, weight=400, color="#1d1d1f",
             lh=None, align="l", spacing=0.0, caps=False):
        """lines: list of str, or list of list[(text, color, weight)] runs."""
        norm = []
        for ln in lines:
            if isinstance(ln, str):
                norm.append([(ln, color, weight)])
            else:
                norm.append([(t, c or color, wt or weight) for (t, c, wt) in ln])
        self.items.append(("text", x, y, w, norm, size, weight, color,
                           lh or size * 1.4, align, spacing, caps))
        return self

    def svg(self, path, x, y, w, h):
        self.items.append(("svg", path, x, y, w, h))
        return self

    # ---- composed blocks --------------------------------------------
    def eyebrow(self, s, y=PAD_Y, color=None):
        self.text(PAD_X, y, 700, [s.upper()], 14, 600,
                  color or ("#2997ff" if self.dark else "#0066cc"), lh=16, spacing=1.6)
        return self

    def head(self, title_lines, y=94, size=44, color=None, rule=True, lh=52):
        self.text(PAD_X, y, 1104, title_lines, size, 600,
                  color or ("#ffffff" if self.dark else "#1d1d1f"), lh=lh, spacing=-1.0)
        if rule:
            ry = y + lh * len(title_lines) + 22
            self.rect(PAD_X, ry, 56, 3, "#2997ff" if self.dark else "#0066cc")
            return ry + 3
        return y + lh * len(title_lines)


def _sv_text(x, y, w, runs, size, lh, align, spacing, caps, i):
    """one line of the QA preview"""
    out = []
    text = "".join(t for t, _, _ in runs)
    if caps:
        text = text.upper()
    anchor = {"l": "start", "c": "middle", "r": "end"}[align]
    tx = x if align == "l" else (x + w / 2.0 if align == "c" else x + w)
    baseline = y + lh * i + size * 0.80
    # runs are rendered as one tspan chain so colors survive the preview
    spans = []
    for t, c, wt in runs:
        if caps:
            t = t.upper()
        fam = EA + ", Arial" if _has_hangul(t) else "Arial"
        spans.append('<tspan fill="%s" font-weight="%d" font-family="%s">%s</tspan>'
                     % (c, wt, fam, _esc(t)))
    out.append('<text x="%.1f" y="%.1f" font-size="%.1f" text-anchor="%s" letter-spacing="%.2f">%s</text>'
               % (tx, baseline, size, anchor, spacing, "".join(spans)))
    return out


def render_svg(slide, asset_dir):
    b = ['<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
         'width="%d" height="%d" viewBox="0 0 %d %d">' % (SLIDE_W, SLIDE_H, SLIDE_W, SLIDE_H)]
    b.append('<rect width="%d" height="%d" fill="%s"/>' % (SLIDE_W, SLIDE_H, slide.bg))
    for it in slide.items:
        kind = it[0]
        if kind == "rect":
            _, x, y, w, h, fill, radius, lineC, lw = it
            st = ' stroke="%s" stroke-width="%s"' % (lineC, lw) if lineC else ""
            b.append('<rect x="%.1f" y="%.1f" width="%.1f" height="%.1f" rx="%.1f" fill="%s"%s/>'
                     % (x, y, w, h, radius, fill or "none", st))
        elif kind == "line":
            _, x1, y1, x2, y2, c, lw = it
            b.append('<line x1="%.1f" y1="%.1f" x2="%.1f" y2="%.1f" stroke="%s" stroke-width="%s"/>'
                     % (x1, y1, x2, y2, c, lw))
        elif kind == "text":
            _, x, y, w, lines, size, weight, color, lh, align, spacing, caps = it
            for i, runs in enumerate(lines):
                b.extend(_sv_text(x, y, w, runs, size, lh, align, spacing, caps, i))
        elif kind == "svg":
            _, path, x, y, w, h = it
            inner = open(os.path.join(asset_dir, path), encoding="utf-8").read()
            inner = re.sub(r"^<svg[^>]*>", "", inner).replace("</svg>", "")
            m = re.search(r'viewBox="0 0 (\d+) (\d+)"',
                          open(os.path.join(asset_dir, path), encoding="utf-8").read())
            sw, sh = (int(m.group(1)), int(m.group(2))) if m else (w, h)
            b.append('<g transform="translate(%.1f,%.1f) scale(%.4f,%.4f)">%s</g>'
                     % (x, y, w / float(sw), h / float(sh), inner))
    if slide.runner:
        b.append('<line x1="0" y1="%d" x2="%d" y2="%d" stroke="%s" stroke-width="1"/>'
                 % (SLIDE_H - RUNNER_H, SLIDE_W, SLIDE_H - RUNNER_H,
                    "rgba(255,255,255,0.16)" if slide.dark else "#e0e0e0"))
        col = "#cccccc" if slide.dark else "#7a7a7a"
        fam = EA + ", Arial"
        b.append('<text x="%d" y="%d" font-size="13" font-family="%s" fill="%s">%s</text>'
                 % (PAD_X, SLIDE_H - RUNNER_H + 34, fam, col, _esc(slide.runner)))
        b.append('<text x="%d" y="%d" font-size="13" font-family="Arial" font-weight="600" '
                 'fill="%s" text-anchor="end">%s</text>'
                 % (SLIDE_W - PAD_X, SLIDE_H - RUNNER_H + 34,
                    "#ffffff" if slide.dark else "#333333", slide.page or ""))
    b.append("</svg>")
    return "".join(b)


def _style_run(run, size, weight, color, spacing):
    f = run.font
    f.size = Pt(size * 0.75)
    f.bold = weight >= 600
    f.color.rgb = _rgb(color)
    f.name = LATIN
    rPr = run._r.get_or_add_rPr()
    for tag, face in (("a:ea", EA), ("a:cs", LATIN)):
        el = rPr.find(qn(tag))
        if el is None:
            el = parse_xml('<%s xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                           'typeface="%s"/>' % (tag, face))
            rPr.append(el)
        else:
            el.set("typeface", face)
    if spacing:
        rPr.set("spc", str(int(round(spacing * 0.75 * 100))))


def _add_svg_picture(slide_obj, pptx_slide, path, x, y, w, h, counter):
    svg_bytes = open(path, "rb").read()
    png = cairosvg.svg2png(bytestring=svg_bytes, output_width=int(w * 2), output_height=int(h * 2),
                           background_color="#00000000")
    pic = pptx_slide.shapes.add_picture(io.BytesIO(png), E(x), E(y), E(w), E(h))
    partname = PackURI("/ppt/media/graphic%d.svg" % counter)
    part = Part(partname, "image/svg+xml", pptx_slide.part.package, svg_bytes)
    rId = pptx_slide.part.relate_to(part, RT.IMAGE)
    blip = pic._element.blipFill.find(qn("a:blip"))
    blip.append(parse_xml(
        '<a:extLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}">'
        '<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'r:embed="%s"/></a:ext></a:extLst>' % rId))
    return pic


def render_pptx(slides, asset_dir, out_path):
    prs = Presentation()
    prs.slide_width = E(SLIDE_W)
    prs.slide_height = E(SLIDE_H)
    blank = prs.slide_layouts[6]
    svg_counter = [0]

    for sl in slides:
        s = prs.slides.add_slide(blank)
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, E(SLIDE_W), E(SLIDE_H))
        bg.fill.solid()
        bg.fill.fore_color.rgb = _rgb(sl.bg)
        bg.line.fill.background()
        bg.shadow.inherit = False

        for it in sl.items:
            kind = it[0]
            if kind == "rect":
                _, x, y, w, h, fill, radius, lineC, lw = it
                shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
                shp = s.shapes.add_shape(shp_type, E(x), E(y), E(w), E(h))
                if radius:
                    adj = min(0.5, float(radius) / min(w, h))
                    shp.adjustments[0] = adj
                if fill:
                    shp.fill.solid()
                    shp.fill.fore_color.rgb = _rgb(fill)
                else:
                    shp.fill.background()
                if lineC:
                    shp.line.color.rgb = _rgb(lineC)
                    shp.line.width = Pt(lw * 0.75)
                else:
                    shp.line.fill.background()
                shp.shadow.inherit = False
            elif kind == "line":
                _, x1, y1, x2, y2, c, lw = it
                cn = s.shapes.add_connector(1, E(x1), E(y1), E(x2), E(y2))
                cn.line.color.rgb = _rgb(c)
                cn.line.width = Pt(lw * 0.75)
            elif kind == "text":
                _, x, y, w, lines, size, weight, color, lh, align, spacing, caps = it
                box = s.shapes.add_textbox(E(x), E(y - size * 0.22), E(w), E(lh * len(lines) + size))
                tf = box.text_frame
                tf.word_wrap = False
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                for i, runs in enumerate(lines):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
                    p.line_spacing = Pt(lh * 0.75)
                    p.space_before = Pt(0)
                    p.space_after = Pt(0)
                    for t, c, wt in runs:
                        r = p.add_run()
                        r.text = t.upper() if caps else t
                        _style_run(r, size, wt, c, spacing)
            elif kind == "svg":
                _, path, x, y, w, h = it
                svg_counter[0] += 1
                _add_svg_picture(sl, s, os.path.join(asset_dir, path), x, y, w, h, svg_counter[0])

        if sl.runner:
            cn = s.shapes.add_connector(1, 0, E(SLIDE_H - RUNNER_H), E(SLIDE_W), E(SLIDE_H - RUNNER_H))
            cn.line.color.rgb = _rgb("#3d3d40" if sl.dark else "#e0e0e0")
            cn.line.width = Pt(0.75)
            box = s.shapes.add_textbox(E(PAD_X), E(SLIDE_H - RUNNER_H + 18), E(700), E(24))
            tf = box.text_frame
            tf.word_wrap = False
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = sl.runner
            _style_run(r, 13, 400, "#cccccc" if sl.dark else "#7a7a7a", 0)
            box2 = s.shapes.add_textbox(E(SLIDE_W - PAD_X - 200), E(SLIDE_H - RUNNER_H + 18), E(200), E(24))
            tf2 = box2.text_frame
            tf2.word_wrap = False
            tf2.margin_left = tf2.margin_right = tf2.margin_top = tf2.margin_bottom = 0
            p2 = tf2.paragraphs[0]
            p2.alignment = PP_ALIGN.RIGHT
            r2 = p2.add_run()
            r2.text = sl.page or ""
            _style_run(r2, 13, 600, "#ffffff" if sl.dark else "#333333", 0)

        if sl.notes:
            s.notes_slide.notes_text_frame.text = sl.notes

    prs.save(out_path)
    return out_path
